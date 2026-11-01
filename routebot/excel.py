"""Excel: чтение входной таблицы, выгрузка готового маршрута (CLI) и отчёта обхода (бот)."""
from __future__ import annotations

import logging
from collections import Counter
from collections.abc import Sequence
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.worksheet import Worksheet

from . import config, links
from .addresses import clean_addr, clean_text, split_contact
from .models import TIME_FMT, Point, RouteState, Status

log = logging.getLogger(__name__)


class InputError(ValueError):
    """Входную таблицу не удалось разобрать — текст можно показать пользователю."""


# ══════════════════════════ Чтение ══════════════════════════

# Классический формат: № | Зона | Точка | Адрес | Владелец | Контакт | Телефон
LEGACY_COLS = dict(num=0, zone=1, name=2, addr=3, owner=4, contact=5, phone=6)

# Ключевые слова заголовков → поле (первое совпадение слева направо).
# «порядок» — чтобы можно было загрузить обратно готовый маршрут или отчёт.
HEADER_KEYS = [
    ("num", ("№", "номер", "порядок")),
    ("zone", ("зона",)),
    ("name", ("точка", "название", "магазин")),
    ("addr", ("адрес",)),
    ("owner", ("владелец",)),
    ("contact", ("контакт",)),
    ("phone", ("телефон",)),
    ("remark", ("примечание",)),
]


def _detect_columns(header: Sequence) -> dict[str, int]:
    cols: dict[str, int] = {}
    for idx, h in enumerate(header):
        h = str(h or "").strip().lower()
        if not h:
            continue
        for field, keys in HEADER_KEYS:
            if field not in cols and any(k in h for k in keys):
                cols[field] = idx
                break
    return cols if {"num", "addr"} <= cols.keys() else dict(LEGACY_COLS)


def _as_int(v) -> int | None:
    if isinstance(v, bool):
        return None
    if isinstance(v, int):
        return v
    if isinstance(v, float) and v.is_integer():
        return int(v)
    if isinstance(v, str) and v.strip().isdigit():
        return int(v.strip())
    return None


def _pick_sheet(wb) -> Worksheet:
    sheets = wb.worksheets
    for ws in sheets:  # сначала лист с «маршрут» в названии
        if "маршрут" in ws.title.lower():
            return ws
    for ws in sheets:
        header = next(ws.iter_rows(max_row=1, values_only=True), ())
        if any("адрес" in str(h or "").lower() for h in header):
            return ws
    return sheets[0]


def read_input(path: Path) -> list[Point]:
    """
    Точки из входной таблицы; первая — СТАРТ. Колонки находятся по заголовкам,
    поэтому подходят и старый формат, и «№ | Посещено | Название | Адрес | Контакт / телефон».
    """
    try:
        wb = load_workbook(path, read_only=True, data_only=True)
    except Exception as e:  # openpyxl бросает разное на битых/чужих файлах
        raise InputError(f"Не удалось открыть Excel-файл: {e}") from e
    try:
        rows = _pick_sheet(wb).iter_rows(values_only=True)
        cols = _detect_columns(next(rows, ()))
        points: list[Point] = []
        for r in rows:
            p = _row_to_point(r, cols, next_num=max((p.num for p in points), default=0) + 1)
            if p:
                points.append(p)
    finally:
        wb.close()  # read_only держит файл открытым — на Windows его потом не перезаписать

    if not points:
        raise InputError(
            "В таблице не найдено ни одной точки. Нужны колонки «№» и «Адрес», "
            "первая строка с данными — СТАРТ."
        )
    return points


def _row_to_point(r: Sequence, cols: dict[str, int], next_num: int) -> Point | None:
    def cell(field: str):
        i = cols.get(field)
        return r[i] if i is not None and i < len(r) else None

    name, addr = clean_text(cell("name")), clean_addr(clean_text(cell("addr")))
    if not name and not addr:
        return None
    raw_num = cell("num")
    num = _as_int(raw_num)
    if num is None:
        if clean_text(raw_num):  # в колонке «№» текст — это заметка/итог, а не точка
            return None
        num = next_num  # номер не проставили — нумеруем сами
    contact, phone = split_contact(clean_text(cell("contact")), clean_text(cell("phone")))
    return Point(
        num=num,
        zone=clean_text(cell("zone")),
        name=name,
        addr=addr,
        owner=clean_text(cell("owner")),
        contact=contact,
        phone=phone,
        remark=clean_text(cell("remark")),
    )


# ══════════════════════════ Оформление ══════════════════════════

FONT = "Arial"
HEADER_FILL = PatternFill("solid", fgColor="1F3A5F")
HEADER_FONT = Font(name=FONT, bold=True, color="FFFFFF")
BODY_FONT = Font(name=FONT, size=11)
LINK_FONT = Font(name=FONT, color="0563C1", underline="single")
STATUS_FILL = {
    Status.VISITED: PatternFill("solid", fgColor="C6EFCE"),
    Status.SKIPPED: PatternFill("solid", fgColor="FFEB9C"),
    Status.CLOSED: PatternFill("solid", fgColor="FFC7CE"),
    Status.PENDING: PatternFill("solid", fgColor="D9E1F2"),
}
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
WRAP = Alignment(vertical="center", wrap_text=True)


def _style_header(row) -> None:
    for c in row:
        if c.value is not None:
            c.font, c.fill, c.alignment = HEADER_FONT, HEADER_FILL, CENTER


def _new_sheet(wb: Workbook, title: str, headers: list[str], widths: list[int]) -> Worksheet:
    ws = wb.create_sheet(title)
    ws.append(headers)
    _style_header(ws[1])
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[ws.cell(1, i).column_letter].width = w
    return ws


def _style_body(ws: Worksheet, center_cols: Sequence[int] = (), skip_cols: Sequence[int] = ()) -> None:
    for row in ws.iter_rows(min_row=2):
        for c in row:
            if c.column in skip_cols:
                continue
            c.font = BODY_FONT
            c.alignment = CENTER if c.column in center_cols else WRAP


def _link(ws: Worksheet, row: int, col: int, url: str, text: str = "Открыть") -> None:
    c = ws.cell(row, col, text)
    c.hyperlink, c.font = url, LINK_FONT


def safe_save(wb: Workbook, path: Path) -> Path:
    """Если файл открыт в Excel (Windows его блокирует), сохраняем рядом с суффиксом."""
    path = Path(path)
    try:
        wb.save(path)
        return path
    except PermissionError:
        for i in range(1, 100):
            alt = path.with_name(f"{path.stem}_{i}{path.suffix}")
            try:
                wb.save(alt)
                log.warning("Файл %s занят, сохранён как %s", path.name, alt.name)
                return alt
            except PermissionError:
                continue
        raise


def _gmaps_chunks_sheet(wb: Workbook, route: Sequence[Point]) -> None:
    """Маршрут в Google Maps частями по 10 точек; следующая часть начинается там, где закончилась."""
    gs = _new_sheet(wb, "Ссылки Google Maps", ["Часть", "Откуда", "Куда", "Точек", "Ссылка"], [8, 26, 26, 8, 20])
    pts = [p for p in route if p.has_coords]
    step = links.GMAPS_MAX_STOPS - 1
    for k, i in enumerate(range(0, max(len(pts) - 1, 0), step), 1):
        chunk = pts[i : i + links.GMAPS_MAX_STOPS]
        gs.append([k, chunk[0].name, chunk[-1].name, len(chunk)])
        _link(gs, gs.max_row, 5, links.gmaps_directions(chunk), "Открыть маршрут")


def _missing_sheet(wb: Workbook, missing: Sequence[Point]) -> None:
    if not missing:
        return
    ms = _new_sheet(
        wb,
        "Не найдены",
        ["№ в исходной таблице", "Точка", "Адрес", "Причина", "Что делать", "Владелец", "Контакт", "Телефон"],
        [12, 24, 40, 44, 60, 24, 22, 18],
    )
    for p in missing:
        ms.append(
            [
                p.num,
                p.name,
                p.addr,
                p.fail_reason or "Не распознан",
                f"Отметьте точку кнопкой «📌 Я тут» на месте или впишите в "
                f"{config.MANUAL_COORDS.name}: {p.num},широта,долгота",
                p.owner,
                p.contact,
                p.phone,
            ]
        )
    _style_body(ms)


# ══════════════════════════ Готовый маршрут (CLI) ══════════════════════════

def write_route(path: Path, route: Sequence[Point], missing: Sequence[Point]) -> Path:
    wb = Workbook()
    wb.remove(wb.active)
    ws = _new_sheet(
        wb,
        "Маршрут",
        ["Порядок", "Точка", "Адрес", "Владелец", "Контакт", "Телефон", "Минут от предыдущей",
         "Всего минут", "Откуда координаты", "Google Maps", "2ГИС"],
        [8, 24, 38, 26, 24, 18, 12, 11, 14, 13, 8],
    )
    total = 0.0
    for i, p in enumerate(route, 1):
        total += p.mins
        ws.append([i, p.name, p.addr, p.owner, p.contact, p.phone, round(p.mins), round(total), p.src])
        gmaps, dgis = links.point_links(p)
        _link(ws, ws.max_row, 10, gmaps)
        _link(ws, ws.max_row, 11, dgis)
    _style_body(ws, center_cols=(1, 7, 8), skip_cols=(10, 11))
    ws.freeze_panes = "C2"
    ws.cell(ws.max_row + 2, 1, f"Время в пути по дорогам без учёта пробок и стоянок: около {round(total)} мин.").font = Font(name=FONT, size=10)

    _gmaps_chunks_sheet(wb, route)
    _missing_sheet(wb, missing)
    return safe_save(wb, path)


# ══════════════════════════ Отчёт обхода (бот) ══════════════════════════

def write_report(path: Path, state: RouteState) -> Path:
    wb = Workbook()
    wb.remove(wb.active)
    _report_route_sheet(wb, state)
    _report_analytics_sheet(wb, state)
    remaining = state.remaining()
    if any(p.has_coords for p in remaining):
        _gmaps_chunks_sheet(wb, remaining)
    _missing_sheet(wb, [p for p in state.points if not p.has_coords])
    return safe_save(wb, path)


def _report_route_sheet(wb: Workbook, state: RouteState) -> None:
    ws = _new_sheet(
        wb,
        "Маршрут",
        ["Порядок", "Точка", "Адрес", "Владелец", "Контакт", "Телефон", "Минут от предыдущей",
         "Всего минут", "Откуда координаты", "Google Maps", "2ГИС", "Статус", "Время визита",
         "Категория / Причина", "Комментарий / Что сказали в магазине", "Фото", "Примечание из таблицы"],
        [8, 26, 38, 24, 22, 18, 12, 11, 14, 13, 10, 14, 18, 24, 52, 30, 30],
    )
    totals = state.cumulative_minutes()
    for p in state.points:
        ws.append(
            [p.order, p.name, p.addr, p.owner, p.contact, p.phone,
             round(p.mins) if p.has_coords else "", round(totals[p.id]) if p.has_coords else "",
             p.src or "не найден", None, None, p.status.label, p.visit_time, p.reason_tag,
             p.notes, ", ".join(p.photos), p.remark]
        )
        r = ws.max_row
        gmaps, dgis = links.point_links(p)
        _link(ws, r, 10, gmaps)
        _link(ws, r, 11, dgis)
        ws.cell(r, 12).fill = STATUS_FILL[p.status]
    _style_body(ws, center_cols=(1, 7, 8, 12, 13), skip_cols=(10, 11))
    ws.freeze_panes = "C2"


def _report_analytics_sheet(wb: Workbook, state: RouteState) -> None:
    an = _new_sheet(wb, "Аналитика оттока", ["Показатель", "Значение"], [44, 18])
    c = state.counts()
    no_coords = sum(1 for p in state.points if not p.has_coords)
    contacted = c[Status.VISITED] + c[Status.CLOSED]
    rows: list[list] = [
        ["Всего точек в маршруте", len(state.points)],
        ["Посещено (успешный визит)", c[Status.VISITED]],
        ["Отложено", c[Status.SKIPPED]],
        ["Магазин закрыт", c[Status.CLOSED]],
        ["Ожидает посещения", c[Status.PENDING]],
        ["Не найдены на карте", no_coords],
        ["Доля успешных визитов (посещено / посещено+закрыто)",
         f"{c[Status.VISITED] / contacted:.0%}" if contacted else "—"],
    ]

    times = sorted(
        t for t in (_parse_time(p.visit_time) for p in state.points if not p.status.is_active) if t
    )
    if times:
        span = (times[-1] - times[0]).total_seconds() / 60
        rows += [
            ["Первая отметка", times[0].strftime(TIME_FMT)],
            ["Последняя отметка", times[-1].strftime(TIME_FMT)],
            ["Время в работе, мин", round(span)],
        ]
        if len(times) > 1:
            rows.append(["В среднем между точками, мин", round(span / (len(times) - 1))])
    for row in rows:
        an.append(row)

    an.append([])
    an.append(["Разбивка по причинам оттока / категориям", "Количество"])
    _style_header(an[an.max_row])
    reasons = Counter(p.reason_tag for p in state.points if p.reason_tag)
    for reason, count in reasons.most_common() or [("(пока нет отмеченных категорий)", 0)]:
        an.append([reason, count])

    zones = Counter(p.zone for p in state.points if p.zone and p.status is Status.VISITED)
    if zones:
        an.append([])
        an.append(["Посещено по зонам", "Количество"])
        _style_header(an[an.max_row])
        for zone, count in zones.most_common():
            an.append([zone, count])


def _parse_time(s: str) -> datetime | None:
    try:
        return datetime.strptime(s, TIME_FMT)
    except (TypeError, ValueError):
        return None
