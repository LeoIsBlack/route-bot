import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook, load_workbook

from routebot.excel import InputError, read_input, write_report, write_route
from routebot.models import Point, RouteState, Status


def make_xlsx(path: Path, title: str, rows: list[list]) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = title
    for r in rows:
        ws.append(r)
    wb.save(path)
    return path


class ReadInputTest(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.dir = Path(self.tmp.name)

    def tearDown(self):
        self.tmp.cleanup()

    def test_classic_format(self):
        path = make_xlsx(self.dir / "a.xlsx", "Маршрут", [
            ["№", "Зона", "Точка", "Адрес (Астана)", "Владелец", "Контакт", "Телефон"],
            [1, "Левый берег", "Старт GS", "ул. Толе би 46", "Манас", "Жамал", "+7 702 222 0723"],
            [2, "Правый", "Второй", "ул. Анет баба 2", None, "Мира", "+7 777 644 1441"],
        ])
        pts = read_input(path)
        self.assertEqual([p.num for p in pts], [1, 2])
        self.assertEqual((pts[0].zone, pts[0].owner, pts[0].phone), ("Левый берег", "Манас", "+7 702 222 0723"))

    def test_new_format_with_combined_contact_and_remark(self):
        path = make_xlsx(self.dir / "b.xlsx", "Маршрут Астана", [
            ["№", "Посещено", "Название точки", "Адрес", "Контакт / телефон", "Примечание"],
            [1, "☐", "Алкохолл GS", "г.Астана, ул.Толе би 46", "Жамал Манас +7 702 222 0723", "звонить заранее"],
            ["2", "☐", "AVO", "г.Астана, ул.Момышулы 2в", "⁠Гульнар +7 771 163 3625", None],
            [3.0, "☐", "Без адреса", None, None, None],
            ["Итого: 3 точки", None, None, None, None, None],
        ])
        pts = read_input(path)
        self.assertEqual([p.num for p in pts], [1, 2, 3])
        self.assertEqual(pts[0].addr, "ул. Толе би 46")
        self.assertEqual((pts[0].contact, pts[0].phone), ("Жамал Манас", "+7 702 222 0723"))
        self.assertEqual(pts[0].remark, "звонить заранее")
        self.assertEqual(pts[1].contact, "Гульнар")
        self.assertEqual(pts[2].addr, "")

    def test_rows_without_number_are_auto_numbered(self):
        path = make_xlsx(self.dir / "c.xlsx", "Лист1", [
            ["№", "Название", "Адрес"],
            [None, "A", "ул. 1"],
            [None, "B", "ул. 2"],
        ])
        self.assertEqual([p.num for p in read_input(path)], [1, 2])

    def test_empty_table_is_input_error(self):
        path = make_xlsx(self.dir / "d.xlsx", "Маршрут", [["№", "Адрес"]])
        with self.assertRaises(InputError):
            read_input(path)

    def test_not_excel_is_input_error(self):
        bad = self.dir / "e.xlsx"
        bad.write_text("not a zip")
        with self.assertRaises(InputError):
            read_input(bad)

    def test_route_output_can_be_uploaded_back(self):
        route = [
            Point(num=1, name="A", addr="ул. 1", lat=51.1, lon=71.4, is_start=True),
            Point(num=2, name="B", addr="ул. 2", lat=51.12, lon=71.42, mins=6.4),
        ]
        out = write_route(self.dir / "route.xlsx", route, [Point(num=3, name="C", addr="?", fail_reason="нет")])
        self.assertEqual([p.name for p in read_input(out)], ["A", "B"])
        wb = load_workbook(out)
        self.assertEqual(wb.sheetnames, ["Маршрут", "Ссылки Google Maps", "Не найдены"])
        self.assertIn("2gis.kz/astana/geo/", wb["Маршрут"].cell(2, 11).hyperlink.target)


class ReportTest(unittest.TestCase):
    def test_report_sheets_and_statuses(self):
        state = RouteState(points=[
            Point(name="A", addr="1", lat=51.1, lon=71.4, status=Status.VISITED, visit_time="2026-11-01 10:00",
                  reason_tag="💸 Цена/Тариф", notes="дорого", photos=["a.jpg"], zone="Левый"),
            Point(name="B", addr="2", lat=51.2, lon=71.5, mins=7, status=Status.CLOSED, visit_time="2026-11-01 10:40"),
            Point(name="C", addr="3", lat=51.3, lon=71.6, mins=9),
            Point(name="D", addr="4"),
        ])
        state.renumber()
        state.ensure_ids()
        with tempfile.TemporaryDirectory() as d:
            wb = load_workbook(write_report(Path(d) / "r.xlsx", state))
            self.assertEqual(wb.sheetnames, ["Маршрут", "Аналитика оттока", "Ссылки Google Maps", "Не найдены"])
            ws = wb["Маршрут"]
            self.assertEqual([ws.cell(r, 12).value for r in range(2, 6)], ["Посетил", "Закрыто", "Ожидает", "Ожидает"])
            self.assertEqual(ws.cell(2, 16).value, "a.jpg")
            analytics = {r[0]: r[1] for r in wb["Аналитика оттока"].iter_rows(values_only=True) if r[0]}
            self.assertEqual(analytics["Время в работе, мин"], 40)
            self.assertEqual(analytics["Доля успешных визитов (посещено / посещено+закрыто)"], "50%")
            self.assertEqual(analytics["💸 Цена/Тариф"], 1)


if __name__ == "__main__":
    unittest.main()
