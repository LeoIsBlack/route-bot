#!/usr/bin/env python3
"""
Telegram-бот для интерактивного полевого обхода магазинов Астаны.

Функции:
  ▶️ Начать / Продолжить обход — карточка текущей точки с inline-кнопками
  📋 Весь маршрут — список всех точек
  📊 Скачать Excel-отчет — итоговый xlsx с аналитикой оттока
  🔄 Сбросить прогресс — очистка состояния
  📤 Отправить xlsx — расчёт маршрута
  📍 Геолокация — пересчёт остатка маршрута от текущего местоположения

Запуск:
    pip install python-telegram-bot openpyxl requests
    set BOT_TOKEN=ваш_токен   (или export BOT_TOKEN=...)
    python tg_bot.py
"""
import asyncio
import json
import logging
import os
import sys
from copy import deepcopy
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook, Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from telegram import (
    Update,
    ReplyKeyboardMarkup,
    InlineKeyboardButton,
    InlineKeyboardMarkup,
    InputFile,
)
from telegram.ext import (
    Application,
    CommandHandler,
    MessageHandler,
    CallbackQueryHandler,
    ConversationHandler,
    filters,
    ContextTypes,
)
from telegram.constants import ParseMode

# ── Настройки ─────────────────────────────────────────────
TOKEN = os.environ.get("BOT_TOKEN", "9093591208:AAHLuair6rfphdfsQKwM8RttPpF6qkAQW3s")
WORK_DIR = Path(__file__).resolve().parent
SRC_FILE = WORK_DIR / "Маршрут_Астана.xlsx"
OUT_FILE = WORK_DIR / "Маршрут_Астана_готовый.xlsx"
SCRIPT = WORK_DIR / "route_planner.py"
STATE_FILE = WORK_DIR / "visit_state.json"
REPORT_FILE = WORK_DIR / "Маршрут_Астана_отчет.xlsx"
PYTHON = sys.executable

logging.basicConfig(
    format="%(asctime)s [%(levelname)s] %(name)s: %(message)s",
    level=logging.INFO,
)
log = logging.getLogger("route-bot")

# Состояния ConversationHandler
WAITING_NOTE = 0

MAIN_KB = ReplyKeyboardMarkup(
    [
        ["▶️ Начать / Продолжить обход", "📋 Весь маршрут"],
        ["📊 Скачать Excel-отчет", "🔄 Сбросить прогресс"],
    ],
    resize_keyboard=True,
)

# Теги причин оттока
REASON_TAGS = {
    "price": "💸 Цена/Тариф",
    "tech": "⚙️ Тех. проблемы/Баги",
    "competitor": "🔄 Ушли к конкуренту",
    "loyal": "🤝 Лояльны / Работают",
    "freetext": "✍️ Сразу текст",
}


# ══════════════════════════════════════════════════════════
# ── Работа с состоянием обхода (visit_state.json) ────────
# ══════════════════════════════════════════════════════════

def _load_state() -> dict:
    """Загружает состояние из JSON-файла."""
    if STATE_FILE.exists():
        with open(STATE_FILE, "r", encoding="utf-8") as f:
            return json.load(f)
    return {}


def _save_state(state: dict):
    """Сохраняет состояние на диск."""
    with open(STATE_FILE, "w", encoding="utf-8") as f:
        json.dump(state, f, ensure_ascii=False, indent=2)


def _init_state_from_route():
    """
    Создаёт/обновляет visit_state.json из Маршрут_Астана_готовый.xlsx.
    Сохраняет данные каждой точки + статус обхода.
    Если точка уже была в state — сохраняет её данные визита.
    """
    points = _read_route_points()
    if not points:
        return {}

    old_state = _load_state()
    old_by_order = {str(p.get("order")): p for p in old_state.get("points", [])}

    new_points = []
    for p in points:
        key = str(p["order"])
        old = old_by_order.get(key, {})
        new_points.append({
            "order": p["order"],
            "name": p["name"],
            "addr": p["addr"],
            "owner": p["owner"],
            "contact": p["contact"],
            "phone": p["phone"],
            "mins": p["mins"],
            "total": p["total"],
            "src": p["src"],
            "gmaps": p["gmaps"],
            "dgis": p["dgis"],
            "lat": p.get("lat"),
            "lon": p.get("lon"),
            # Статус визита (сохраняем старый если был)
            "status": old.get("status", "pending"),
            "visit_time": old.get("visit_time", ""),
            "reason_tag": old.get("reason_tag", ""),
            "notes": old.get("notes", ""),
        })

    state = {"points": new_points}
    _save_state(state)
    return state


def _get_point_by_order(state: dict, order: int) -> dict | None:
    """Находит точку по порядковому номеру."""
    for p in state.get("points", []):
        if p["order"] == order:
            return p
    return None


def _first_pending(state: dict) -> dict | None:
    """Первая точка со статусом 'pending'."""
    for p in state.get("points", []):
        if p["status"] == "pending":
            return p
    return None


def _progress_summary(state: dict) -> str:
    """Строка со сводкой прогресса."""
    points = state.get("points", [])
    total = len(points)
    visited = sum(1 for p in points if p["status"] == "visited")
    skipped = sum(1 for p in points if p["status"] == "skipped")
    closed = sum(1 for p in points if p["status"] == "closed")
    pending = sum(1 for p in points if p["status"] == "pending")
    return (
        f"📊 *Прогресс:* {visited} из {total} посещено\n"
        f"   ✅ Посещено: {visited} | ⏭ Отложено: {skipped} | "
        f"🔒 Закрыто: {closed} | ⏳ Ожидает: {pending}"
    )


# ══════════════════════════════════════════════════════════
# ── Чтение xlsx ──────────────────────────────────────────
# ══════════════════════════════════════════════════════════

def _read_route_points():
    """Читает точки из готового xlsx. Возвращает список dict."""
    if not OUT_FILE.exists():
        return []

    wb = load_workbook(OUT_FILE)
    ws = wb["Маршрут"]

    points = []
    for row in ws.iter_rows(min_row=2, max_col=11):
        val = row[0].value
        if val is None or not isinstance(val, int):
            continue
        gmaps = row[9].hyperlink.target if row[9].hyperlink else None
        dgis = row[10].hyperlink.target if row[10].hyperlink else None

        # Извлекаем координаты из ссылки Google Maps
        lat, lon = None, None
        if gmaps and "query=" in gmaps:
            try:
                coords_str = gmaps.split("query=")[1].split("&")[0]
                lat, lon = [float(x) for x in coords_str.split(",")]
            except (ValueError, IndexError):
                pass

        points.append(dict(
            order=val,
            name=row[1].value or "",
            addr=row[2].value or "",
            owner=row[3].value or "",
            contact=row[4].value or "",
            phone=row[5].value or "",
            mins=row[6].value or 0,
            total=row[7].value or 0,
            src=row[8].value or "",
            gmaps=gmaps,
            dgis=dgis,
            lat=lat,
            lon=lon,
        ))

    return points


# ══════════════════════════════════════════════════════════
# ── Форматирование ───────────────────────────────────────
# ══════════════════════════════════════════════════════════

def _esc(text: str) -> str:
    """Экранирует спец-символы Markdown v1 в тексте."""
    for ch in ("_", "*", "`", "["):
        text = text.replace(ch, "\\" + ch)
    return text


STATUS_EMOJI = {
    "pending": "⏳",
    "visited": "✅",
    "skipped": "⏭",
    "closed": "🔒",
}


def _card_text(p: dict) -> str:
    """Формирует текст карточки точки для Telegram."""
    status_em = STATUS_EMOJI.get(p["status"], "")
    lines = [
        f"*{p['order']}.* {_esc(p['name'])} {status_em}",
        f"📍 {_esc(p['addr'])}",
    ]
    if p.get("owner"):
        lines.append(f"👤 Владелец: {_esc(p['owner'])}")
    if p.get("contact"):
        lines.append(f"🧑‍💼 Контакт: {_esc(p['contact'])}")
    if p.get("phone"):
        lines.append(f"📞 {_esc(p['phone'])}")

    if p["order"] == 1:
        lines.append("🏁 СТАРТ")
    else:
        lines.append(f"⏱ +{p['mins']} мин от предыдущей (всего {p['total']})")

    # Ссылки
    links = []
    if p.get("gmaps"):
        links.append(f"[Google Maps]({p['gmaps']})")
    if p.get("dgis"):
        links.append(f"[2ГИС]({p['dgis']})")
    if links:
        lines.append("🔗 " + " · ".join(links))

    # Заметка если есть
    if p.get("notes"):
        lines.append(f"\n💬 _{_esc(p['notes'])}_")
    if p.get("reason_tag"):
        lines.append(f"🏷 {p['reason_tag']}")
    if p.get("visit_time"):
        lines.append(f"🕐 {p['visit_time']}")

    return "\n".join(lines)


def _card_keyboard(p: dict) -> InlineKeyboardMarkup:
    """Inline-клавиатура под карточкой точки."""
    order = p["order"]
    state = _load_state()
    total = len(state.get("points", []))

    buttons = []

    if p["status"] == "visited":
        buttons.append([
            InlineKeyboardButton("✍️ Дополнить заметку", callback_data=f"visit:{order}")
        ])
    else:
        buttons.append([
            InlineKeyboardButton("✅ Записать итог / Заметку", callback_data=f"visit:{order}")
        ])
    buttons.append([
        InlineKeyboardButton("⏭ Отложить на потом", callback_data=f"skip:{order}"),
        InlineKeyboardButton("🔒 Магазин закрыт", callback_data=f"closed:{order}"),
    ])

    # Навигация
    nav = []
    if order > 1:
        nav.append(InlineKeyboardButton("⬅️ Пред.", callback_data=f"nav:{order - 1}"))
    if order < total:
        nav.append(InlineKeyboardButton("След. ➡️", callback_data=f"nav:{order + 1}"))
    if nav:
        buttons.append(nav)

    return InlineKeyboardMarkup(buttons)


# ══════════════════════════════════════════════════════════
# ── Генерация Excel-отчёта ───────────────────────────────
# ══════════════════════════════════════════════════════════

def export_report_xlsx() -> Path:
    """
    Берёт Маршрут_Астана_готовый.xlsx, добавляет столбцы отчёта
    и лист аналитики оттока. Возвращает путь к файлу.
    """
    state = _load_state()
    points_state = {p["order"]: p for p in state.get("points", [])}

    wb = load_workbook(OUT_FILE)
    ws = wb["Маршрут"]

    F = "Arial"
    hf = PatternFill("solid", fgColor="1F3A5F")

    # Заголовки новых столбцов (L=12, M=13, N=14, O=15)
    new_headers = ["Статус", "Время визита", "Категория / Причина",
                   "Комментарий / Что сказали в магазине"]
    for col_idx, header in enumerate(new_headers, 12):
        cell = ws.cell(1, col_idx, header)
        cell.font = Font(name=F, bold=True, color="FFFFFF")
        cell.fill = hf
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    # Ширины новых столбцов
    ws.column_dimensions["L"].width = 14
    ws.column_dimensions["M"].width = 18
    ws.column_dimensions["N"].width = 22
    ws.column_dimensions["O"].width = 50

    # Цвета заливки
    fill_visited = PatternFill("solid", fgColor="C6EFCE")   # зелёный
    fill_skipped = PatternFill("solid", fgColor="FFEB9C")   # жёлтый
    fill_closed = PatternFill("solid", fgColor="FFC7CE")    # красный
    fill_pending = PatternFill("solid", fgColor="D9E1F2")   # голубой

    STATUS_RU = {
        "visited": "Посетил",
        "skipped": "Пропущено",
        "closed": "Закрыто",
        "pending": "Ожидает",
    }
    STATUS_FILL = {
        "visited": fill_visited,
        "skipped": fill_skipped,
        "closed": fill_closed,
        "pending": fill_pending,
    }

    for row in ws.iter_rows(min_row=2, max_col=1):
        order_val = row[0].value
        if order_val is None or not isinstance(order_val, int):
            continue
        r = row[0].row
        ps = points_state.get(order_val, {})
        status = ps.get("status", "pending")

        # Столбец L — Статус
        cell_status = ws.cell(r, 12, STATUS_RU.get(status, status))
        cell_status.fill = STATUS_FILL.get(status, fill_pending)
        cell_status.font = Font(name=F, size=11)
        cell_status.alignment = Alignment(horizontal="center", vertical="center")

        # Столбец M — Время визита
        cell_time = ws.cell(r, 13, ps.get("visit_time", ""))
        cell_time.font = Font(name=F, size=11)
        cell_time.alignment = Alignment(horizontal="center", vertical="center")

        # Столбец N — Категория
        cell_reason = ws.cell(r, 14, ps.get("reason_tag", ""))
        cell_reason.font = Font(name=F, size=11)
        cell_reason.alignment = Alignment(vertical="center", wrap_text=True)

        # Столбец O — Комментарий
        cell_notes = ws.cell(r, 15, ps.get("notes", ""))
        cell_notes.font = Font(name=F, size=11)
        cell_notes.alignment = Alignment(vertical="center", wrap_text=True)

    # ── Лист «Аналитика оттока» ──
    if "Аналитика оттока" in wb.sheetnames:
        del wb["Аналитика оттока"]
    an = wb.create_sheet("Аналитика оттока")

    all_points = state.get("points", [])
    total = len(all_points)
    visited = [p for p in all_points if p["status"] == "visited"]
    skipped = [p for p in all_points if p["status"] == "skipped"]
    closed = [p for p in all_points if p["status"] == "closed"]
    pending = [p for p in all_points if p["status"] == "pending"]

    an.append(["Показатель", "Количество"])
    an.append(["Всего точек", total])
    an.append(["Посещено", len(visited)])
    an.append(["Отложено", len(skipped)])
    an.append(["Закрыто", len(closed)])
    an.append(["Ожидает", len(pending)])
    an.append([])
    an.append(["Разбивка по причинам оттока", "Количество"])

    # Считаем причины
    reason_counts = {}
    for p in all_points:
        tag = p.get("reason_tag", "")
        if tag:
            reason_counts[tag] = reason_counts.get(tag, 0) + 1
    if reason_counts:
        for reason, count in sorted(reason_counts.items(), key=lambda x: -x[1]):
            an.append([reason, count])
    else:
        an.append(["(нет данных)", 0])

    # Стили
    for c in an[1]:
        c.font = Font(name=F, bold=True, color="FFFFFF")
        c.fill = hf
        c.alignment = Alignment(horizontal="center", vertical="center")
    row8 = an[8]
    if row8:
        for c in row8:
            if c.value:
                c.font = Font(name=F, bold=True, color="FFFFFF")
                c.fill = hf
                c.alignment = Alignment(horizontal="center", vertical="center")

    an.column_dimensions["A"].width = 35
    an.column_dimensions["B"].width = 14

    for row in an.iter_rows(min_row=2, max_col=2):
        for c in row:
            if c.font == Font():
                c.font = Font(name=F, size=11)

    # Сохранение
    try:
        wb.save(str(REPORT_FILE))
    except PermissionError:
        alt = WORK_DIR / f"Маршрут_Астана_отчет_{datetime.now():%H%M%S}.xlsx"
        wb.save(str(alt))
        return alt
    return REPORT_FILE


# ══════════════════════════════════════════════════════════
# ── Команды и хэндлеры бота ──────────────────────────────
# ══════════════════════════════════════════════════════════

async def cmd_start(update: Update, _ctx: ContextTypes.DEFAULT_TYPE):
    await update.message.reply_text(
        "👋 *Бот полевого обхода магазинов Астаны*\n\n"
        "▶️ *Начать / Продолжить обход* — карточка текущей точки\n"
        "📋 *Весь маршрут* — список всех точек\n"
        "📊 *Скачать Excel-отчет* — итоговый файл\n"
        "🔄 *Сбросить прогресс* — начать сначала\n"
        "📤 Отправьте *.xlsx* — пересчитаю маршрут\n"
        "📍 Отправьте *геолокацию* — пересчёт от вашей точки\n\n"
        "Все ссылки Google Maps и 2ГИС кликабельны 👆",
        parse_mode=ParseMode.MARKDOWN,
        reply_markup=MAIN_KB,
    )


# ── Показать весь маршрут ────────────────────────────────

async def cmd_show_all(update: Update, _ctx: ContextTypes.DEFAULT_TYPE):
    state = _load_state()
    points = state.get("points", [])
    if not points:
        await update.message.reply_text(
            "📤 Сначала отправьте xlsx-файл с адресами.\n"
            "Я рассчитаю маршрут и покажу его здесь.",
            reply_markup=MAIN_KB,
        )
        return

    msgs, buf = [], ["🗺 *Маршрут по Астане*\n"]
    for p in points:
        em = STATUS_EMOJI.get(p["status"], "")
        links = []
        if p.get("gmaps"):
            links.append(f'[GMaps]({p["gmaps"]})')
        if p.get("dgis"):
            links.append(f'[2ГИС]({p["dgis"]})')
        lnk = " · ".join(links)

        if p["order"] == 1:
            time_s = "🏁 СТАРТ"
        else:
            time_s = f"⏱ +{p['mins']} мин  (всего {p['total']})"

        block = (
            f"*{p['order']}.* {_esc(p['name'])} {em}\n"
            f"   📍 {_esc(p['addr'])}\n"
            f"   {time_s}\n"
            f"   {lnk}\n"
        )

        joined = "\n".join(buf) + "\n" + block
        if len(joined) > 3800:
            msgs.append("\n".join(buf))
            buf = ["🗺 *Продолжение*\n"]
        buf.append(block)

    total_time = points[-1].get("total", 0) if points else 0
    buf.append(f"\n⏱ *Общее время в пути: ≈ {total_time} мин*")
    buf.append("\n" + _progress_summary(state))
    msgs.append("\n".join(buf))

    for msg in msgs:
        await update.message.reply_text(
            msg,
            parse_mode=ParseMode.MARKDOWN,
            disable_web_page_preview=True,
        )


# ── Начать/Продолжить обход ──────────────────────────────

async def _send_point_card(target, state: dict, order: int | None = None):
    """
    Отправляет карточку точки. target — Message или CallbackQuery.
    Если order не указан, берёт первую pending.
    """
    if order is not None:
        p = _get_point_by_order(state, order)
    else:
        p = _first_pending(state)

    if p is None:
        text = "🎉 *Все точки пройдены\\!*\n\n" + _progress_summary(state)
        if hasattr(target, "message"):
            await target.message.reply_text(
                text, parse_mode=ParseMode.MARKDOWN, reply_markup=MAIN_KB
            )
        else:
            await target.reply_text(
                text, parse_mode=ParseMode.MARKDOWN, reply_markup=MAIN_KB
            )
        return None

    text = _card_text(p)
    kb = _card_keyboard(p)

    if hasattr(target, "message") and target.message:
        # CallbackQuery
        await target.message.reply_text(
            text, parse_mode=ParseMode.MARKDOWN,
            disable_web_page_preview=True, reply_markup=kb
        )
    else:
        await target.reply_text(
            text, parse_mode=ParseMode.MARKDOWN,
            disable_web_page_preview=True, reply_markup=kb
        )
    return p


async def cmd_start_walk(update: Update, _ctx: ContextTypes.DEFAULT_TYPE):
    state = _load_state()
    if not state.get("points"):
        await update.message.reply_text(
            "📤 Сначала отправьте xlsx-файл с адресами.",
            reply_markup=MAIN_KB,
        )
        return

    result = await _send_point_card(update.message, state)
    if result is None:
        # Все пройдены — предлагаем скачать отчёт
        await _auto_send_report(update.message)


async def _auto_send_report(message):
    """Автоматическая отправка отчёта при завершении обхода."""
    try:
        report_path = export_report_xlsx()
        with open(report_path, "rb") as f:
            await message.reply_document(
                document=InputFile(f, filename="Маршрут_Астана_отчет.xlsx"),
                caption="📊 Итоговый отчёт обхода — автоматически сформирован",
            )
    except Exception as e:
        log.error(f"Ошибка автоотчёта: {e}")


# ── Callback-кнопки навигации, skip, closed ──────────────

async def cb_navigation(update: Update, _ctx: ContextTypes.DEFAULT_TYPE):
    """Обработка nav:{order} — переключение между точками."""
    query = update.callback_query
    await query.answer()
    _, order_str = query.data.split(":", 1)
    order = int(order_str)
    state = _load_state()
    await _send_point_card(query, state, order)


async def cb_skip(update: Update, _ctx: ContextTypes.DEFAULT_TYPE):
    """Обработка skip:{order} — отложить на потом."""
    query = update.callback_query
    await query.answer("⏭ Отложено")
    _, order_str = query.data.split(":", 1)
    order = int(order_str)

    state = _load_state()
    p = _get_point_by_order(state, order)
    if p and p["status"] in ("pending", "skipped"):
        p["status"] = "skipped"
        _save_state(state)

    # Показать следующую pending
    await _send_point_card(query, state)


async def cb_closed(update: Update, _ctx: ContextTypes.DEFAULT_TYPE):
    """Обработка closed:{order} — магазин закрыт."""
    query = update.callback_query
    await query.answer("🔒 Отмечено как закрытый")
    _, order_str = query.data.split(":", 1)
    order = int(order_str)

    state = _load_state()
    p = _get_point_by_order(state, order)
    if p:
        p["status"] = "closed"
        p["visit_time"] = datetime.now().strftime("%Y-%m-%d %H:%M")
        _save_state(state)

    summary = _progress_summary(state)
    await query.message.reply_text(summary, parse_mode=ParseMode.MARKDOWN)

    # Проверяем завершение
    if _first_pending(state) is None:
        await query.message.reply_text(
            "🎉 *Все точки пройдены\\!*", parse_mode=ParseMode.MARKDOWN,
            reply_markup=MAIN_KB,
        )
        await _auto_send_report(query.message)
    else:
        await _send_point_card(query, state)


# ── Диалог визита (ConversationHandler) ──────────────────

async def cb_visit_start(update: Update, ctx: ContextTypes.DEFAULT_TYPE):
    """
    visit:{order} — начало диалога фиксации визита.
    Показывает inline-кнопки с тегами причин.
    """
    query = update.callback_query
    await query.answer()
    _, order_str = query.data.split(":", 1)
    order = int(order_str)

    ctx.user_data["visit_order"] = order

    state = _load_state()
    p = _get_point_by_order(state, order)
    is_append = p and p["status"] == "visited"

    prompt = "✍️ Дополните заметку — выберите тег:" if is_append else "Выберите категорию:"

    kb = InlineKeyboardMarkup([
        [InlineKeyboardButton("💸 Цена/Тариф", callback_data="tag:price")],
        [InlineKeyboardButton("⚙️ Тех. проблемы/Баги", callback_data="tag:tech")],
        [InlineKeyboardButton("🔄 Ушли к конкуренту", callback_data="tag:competitor")],
        [InlineKeyboardButton("🤝 Лояльны / Работают", callback_data="tag:loyal")],
        [InlineKeyboardButton("✍️ Сразу текст", callback_data="tag:freetext")],
    ])
    await query.message.reply_text(prompt, reply_markup=kb)
    return WAITING_NOTE


async def cb_tag_selected(update: Update, ctx: ContextTypes.DEFAULT_TYPE):
    """Выбор тега — переводим в WAITING_NOTE."""
    query = update.callback_query
    await query.answer()
    _, tag_key = query.data.split(":", 1)

    tag_label = REASON_TAGS.get(tag_key, tag_key)
    ctx.user_data["visit_tag"] = tag_label

    if tag_key == "freetext":
        await query.message.reply_text(
            "✍️ Напишите, что сказали в магазине:"
        )
    else:
        await query.message.reply_text(
            f"Тег: *{tag_label}*\n\n"
            "Напишите, что сказали в магазине "
            "(или отправьте *\\-* если без комментариев):",
            parse_mode=ParseMode.MARKDOWN,
        )
    return WAITING_NOTE


async def handle_note_text(update: Update, ctx: ContextTypes.DEFAULT_TYPE):
    """Получен текст заметки — сохраняем и переходим к следующей точке."""
    order = ctx.user_data.get("visit_order")
    tag = ctx.user_data.get("visit_tag", "")
    note_text = update.message.text.strip()

    if note_text == "-":
        note_text = ""

    state = _load_state()
    p = _get_point_by_order(state, order)

    if p:
        is_append = p["status"] == "visited"

        if is_append and note_text:
            # Дописываем к существующей заметке
            old_notes = p.get("notes", "")
            if old_notes:
                p["notes"] = old_notes + "; " + note_text
            else:
                p["notes"] = note_text
            # Обновляем тег если выбран новый
            if tag and tag != "✍️ Сразу текст":
                p["reason_tag"] = tag
        else:
            p["status"] = "visited"
            p["visit_time"] = datetime.now().strftime("%Y-%m-%d %H:%M")
            if tag and tag != "✍️ Сразу текст":
                p["reason_tag"] = tag
            p["notes"] = note_text

        _save_state(state)

    await update.message.reply_text(
        f"✅ Точка *{order}* сохранена\\!\n\n" + _progress_summary(state),
        parse_mode=ParseMode.MARKDOWN,
        reply_markup=MAIN_KB,
    )

    # Проверяем завершение обхода
    next_p = _first_pending(state)
    if next_p is None:
        all_done = all(
            pt["status"] != "pending" for pt in state.get("points", [])
        )
        if all_done:
            await update.message.reply_text(
                "🎉 *Все точки маршрута пройдены\\!*",
                parse_mode=ParseMode.MARKDOWN,
            )
            await _auto_send_report(update.message)
    else:
        await _send_point_card(update.message, state)

    # Очищаем
    ctx.user_data.pop("visit_order", None)
    ctx.user_data.pop("visit_tag", None)
    return ConversationHandler.END


async def cancel_visit(update: Update, ctx: ContextTypes.DEFAULT_TYPE):
    """Отмена диалога визита."""
    ctx.user_data.pop("visit_order", None)
    ctx.user_data.pop("visit_tag", None)
    await update.message.reply_text("❌ Отменено.", reply_markup=MAIN_KB)
    return ConversationHandler.END


# ── Скачать Excel-отчет ──────────────────────────────────

async def cmd_download_report(update: Update, _ctx: ContextTypes.DEFAULT_TYPE):
    state = _load_state()
    if not state.get("points"):
        await update.message.reply_text(
            "📤 Сначала загрузите xlsx и рассчитайте маршрут.",
            reply_markup=MAIN_KB,
        )
        return

    wait_msg = await update.message.reply_text("⏳ Формирую отчёт…")

    try:
        report_path = export_report_xlsx()
        with open(report_path, "rb") as f:
            await update.message.reply_document(
                document=InputFile(f, filename="Маршрут_Астана_отчет.xlsx"),
                caption="📊 Итоговый Excel-отчёт обхода",
            )
        await wait_msg.delete()
    except Exception as e:
        await wait_msg.edit_text(f"❌ Ошибка: {e}")


# ── Сбросить прогресс ────────────────────────────────────

async def cmd_reset(update: Update, _ctx: ContextTypes.DEFAULT_TYPE):
    state = _load_state()
    if not state.get("points"):
        await update.message.reply_text("Нечего сбрасывать.", reply_markup=MAIN_KB)
        return

    for p in state["points"]:
        p["status"] = "pending"
        p["visit_time"] = ""
        p["reason_tag"] = ""
        p["notes"] = ""
    _save_state(state)

    await update.message.reply_text(
        "🔄 Прогресс сброшен. Все точки снова в статусе «Ожидает».",
        reply_markup=MAIN_KB,
    )


# ── Обработка xlsx-файла ─────────────────────────────────

async def handle_document(update: Update, ctx: ContextTypes.DEFAULT_TYPE):
    doc = update.message.document
    if doc.mime_type not in (
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        "application/vnd.ms-excel",
    ) and not (doc.file_name or "").lower().endswith(".xlsx"):
        await update.message.reply_text(
            "❌ Отправьте файл в формате *.xlsx*",
            parse_mode=ParseMode.MARKDOWN,
        )
        return

    status = await update.message.reply_text("⏳ Загружаю файл…", reply_markup=MAIN_KB)

    try:
        tg_file = await ctx.bot.get_file(doc.file_id)
        await tg_file.download_to_drive(str(SRC_FILE))
    except Exception as e:
        await status.edit_text(f"❌ Не удалось скачать файл: {e}")
        return

    await status.edit_text(
        "🔄 Считаю маршрут…\nПоиск адресов ≈ 1 сек/адрес, подождите ⏳"
    )

    try:
        proc = await asyncio.create_subprocess_exec(
            PYTHON, str(SCRIPT),
            cwd=str(WORK_DIR),
            stdout=asyncio.subprocess.PIPE,
            stderr=asyncio.subprocess.STDOUT,
        )
        stdout, _ = await asyncio.wait_for(proc.communicate(), timeout=600)
        output = stdout.decode("utf-8", errors="replace")
    except asyncio.TimeoutError:
        await status.edit_text("❌ Таймаут: расчёт занял больше 10 минут.")
        return
    except Exception as e:
        await status.edit_text(f"❌ Ошибка запуска скрипта: {e}")
        return

    if proc.returncode != 0:
        snippet = output[:3000] if output else "неизвестная ошибка"
        await status.edit_text(
            f"❌ Ошибка при расчёте:\n```\n{snippet}\n```",
            parse_mode=ParseMode.MARKDOWN,
        )
        return

    # Инициализируем visit_state.json
    _init_state_from_route()
    ctx.user_data["has_route"] = True
    await status.edit_text("✅ Маршрут готов!")

    # Отправляем файл
    with open(OUT_FILE, "rb") as f:
        await update.message.reply_document(
            document=InputFile(f, filename="Маршрут_Астана_готовый.xlsx"),
            caption="📎 Файл маршрута",
        )

    # Лог расчёта
    if output.strip():
        if len(output) > 3800:
            output = output[:3800] + "\n…"
        await update.message.reply_text(
            f"📝 *Лог расчёта:*\n```\n{output}\n```",
            parse_mode=ParseMode.MARKDOWN,
        )

    state = _load_state()
    await update.message.reply_text(
        _progress_summary(state) + "\n\nНажмите *▶️ Начать / Продолжить обход*",
        parse_mode=ParseMode.MARKDOWN,
        reply_markup=MAIN_KB,
    )


# ── Обработка геолокации — пересчёт маршрута ─────────────

async def handle_location(update: Update, _ctx: ContextTypes.DEFAULT_TYPE):
    """
    При получении геолокации: пересчитывает порядок оставшихся
    (pending + skipped) точек от текущей позиции через OSRM.
    """
    loc = update.message.location
    lat, lon = loc.latitude, loc.longitude

    state = _load_state()
    points = state.get("points", [])
    if not points:
        await update.message.reply_text(
            "📤 Сначала загрузите xlsx и рассчитайте маршрут.",
            reply_markup=MAIN_KB,
        )
        return

    # Отфильтровываем remaining (pending + skipped) с координатами
    remaining = [
        p for p in points
        if p["status"] in ("pending", "skipped") and p.get("lat") and p.get("lon")
    ]

    if not remaining:
        await update.message.reply_text(
            "✅ Все точки уже обработаны, пересчитывать нечего.",
            reply_markup=MAIN_KB,
        )
        return

    wait_msg = await update.message.reply_text(
        f"📍 Ваши координаты: {lat:.6f}, {lon:.6f}\n"
        f"🔄 Пересчитываю маршрут для {len(remaining)} оставшихся точек…"
    )

    try:
        # Импортируем функции из route_planner
        sys.path.insert(0, str(WORK_DIR))
        from route_planner import reroute_remaining

        # Подготавливаем точки для route_planner (нужен формат с xy)
        rp_points = []
        for p in remaining:
            rp_points.append({
                "order": p["order"],
                "name": p["name"],
                "addr": p["addr"],
                "owner": p["owner"],
                "contact": p["contact"],
                "phone": p["phone"],
                "xy": (p["lat"], p["lon"]),
                "src": p["src"],
                # Сохраняем данные визита
                "_status": p["status"],
                "_visit_time": p.get("visit_time", ""),
                "_reason_tag": p.get("reason_tag", ""),
                "_notes": p.get("notes", ""),
                "_gmaps": p.get("gmaps", ""),
                "_dgis": p.get("dgis", ""),
            })

        ordered, minutes = reroute_remaining(lat, lon, rp_points)

    except Exception as e:
        await wait_msg.edit_text(f"❌ Ошибка пересчёта: {e}")
        return

    # Сохраняем visited/closed точки без изменений,
    # remaining — в новом порядке с пересчитанным временем
    visited_closed = [p for p in points if p["status"] in ("visited", "closed")]

    new_points = []
    # Сначала — уже обработанные (сохраняем их порядок)
    for p in visited_closed:
        new_points.append(p)

    # Потом — пересортированные remaining
    cumulative = 0
    for idx, (rp, m) in enumerate(zip(ordered, minutes)):
        cumulative += m
        new_order = len(visited_closed) + idx + 1
        new_points.append({
            "order": new_order,
            "name": rp["name"],
            "addr": rp["addr"],
            "owner": rp["owner"],
            "contact": rp["contact"],
            "phone": rp["phone"],
            "mins": round(m),
            "total": round(cumulative),
            "src": rp["src"],
            "gmaps": rp.get("_gmaps", ""),
            "dgis": rp.get("_dgis", ""),
            "lat": rp["xy"][0],
            "lon": rp["xy"][1],
            "status": rp.get("_status", "pending"),
            "visit_time": rp.get("_visit_time", ""),
            "reason_tag": rp.get("_reason_tag", ""),
            "notes": rp.get("_notes", ""),
        })

    # Пересчитываем порядок для всех
    for i, p in enumerate(new_points, 1):
        p["order"] = i

    state["points"] = new_points
    _save_state(state)

    await wait_msg.edit_text(
        f"✅ Маршрут пересчитан от вашей позиции!\n"
        f"📍 {lat:.4f}, {lon:.4f}\n"
        f"Осталось точек: {len(ordered)}\n\n"
        + _progress_summary(state),
        parse_mode=ParseMode.MARKDOWN,
    )

    # Показываем первую pending точку
    await _send_point_card(update.message, state)


# ══════════════════════════════════════════════════════════
# ── Запуск ───────────────────────────────────────────────
# ══════════════════════════════════════════════════════════

def main():
    app = Application.builder().token(TOKEN).build()

    # ConversationHandler для диалога визита
    visit_conv = ConversationHandler(
        entry_points=[
            CallbackQueryHandler(cb_visit_start, pattern=r"^visit:\d+$"),
        ],
        states={
            WAITING_NOTE: [
                CallbackQueryHandler(cb_tag_selected, pattern=r"^tag:"),
                MessageHandler(filters.TEXT & ~filters.COMMAND, handle_note_text),
            ],
        },
        fallbacks=[
            CommandHandler("cancel", cancel_visit),
            MessageHandler(filters.Regex(r"^🔄"), cancel_visit),
        ],
        per_message=False,
    )

    app.add_handler(visit_conv)

    app.add_handler(CommandHandler("start", cmd_start))
    app.add_handler(CommandHandler("route", cmd_show_all))

    # Reply-клавиатура
    app.add_handler(MessageHandler(
        filters.Regex(r"^▶️"), cmd_start_walk
    ))
    app.add_handler(MessageHandler(
        filters.Regex(r"^📋"), cmd_show_all
    ))
    app.add_handler(MessageHandler(
        filters.Regex(r"^📊"), cmd_download_report
    ))
    app.add_handler(MessageHandler(
        filters.Regex(r"^🔄"), cmd_reset
    ))

    # Inline callback'и (навигация, skip, closed)
    app.add_handler(CallbackQueryHandler(cb_navigation, pattern=r"^nav:\d+$"))
    app.add_handler(CallbackQueryHandler(cb_skip, pattern=r"^skip:\d+$"))
    app.add_handler(CallbackQueryHandler(cb_closed, pattern=r"^closed:\d+$"))

    # Документы (xlsx)
    app.add_handler(MessageHandler(filters.Document.ALL, handle_document))

    # Геолокация
    app.add_handler(MessageHandler(filters.LOCATION, handle_location))

    log.info("🤖 Бот запущен, ожидаю сообщения…")
    app.run_polling(drop_pending_updates=True)


if __name__ == "__main__":
    main()
